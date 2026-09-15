from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parent / "Bao_cao_trien_khai_KPI_Admin_Web.docx"

NAVY = "17365D"
BLUE = "2F5597"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "D9E1F2"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="B7C9E2", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=45, start=90, bottom=45, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for child in (begin, instr, separate, text, end):
        run._r.append(child)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{label}: ")
    r.bold = True
    p.add_run(value)
    return p


def set_table_geometry(table, widths):
    column_widths = [round(width * 1440) for width in widths]
    table_width = sum(column_widths)
    table_indent = 90
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(table_width))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(table_indent))

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in column_widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)


def add_table(doc, headers, rows, widths=None, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        set_table_geometry(table, widths)
    header = table.rows[0]
    repeat_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        cell.text = value
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
        set_cell_border(cell)
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[idx])
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if len(table.rows) % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(8.2)
                    run.font.color.rgb = TEXT
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = TEXT
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.12

for level, size, color in ((1, 16, "000000"), (2, 12, "000000"), (3, 10.5, "000000")):
    style = styles[f"Heading {level}"]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    styles[style_name].font.name = "Aptos"
    styles[style_name].font.size = Pt(10)
    styles[style_name].paragraph_format.space_after = Pt(3)

# Header and footer
header = section.header
hp = header.paragraphs[0]
hp.text = "TRUNG TÂM THAM VẤN TÂM LÝ HỌC ĐƯỜNG  |  BÁO CÁO TRIỂN KHAI"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in hp.runs:
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("Nội bộ · Không chứa dữ liệu định danh học sinh  |  Trang ")
add_field(fp, "PAGE")
fp.add_run(" / ")
add_field(fp, "NUMPAGES")
for run in fp.runs:
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(48)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("BÁO CÁO TRIỂN KHAI")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(15)
r.font.color.rgb = RGBColor(0, 0, 0)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(10)
r = p.add_run("MÔ HÌNH ĐÁNH GIÁ KPI\nVÀ CỔNG WEB MỘT TÀI KHOẢN ADMIN")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(25)
r.font.color.rgb = RGBColor(0, 0, 0)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(28)
r = p.add_run("Dự án Digital Twin – Trung tâm Tham vấn Tâm lý Học đường")
r.font.size = Pt(12)
r.font.color.rgb = MUTED

meta = add_table(
    doc,
    ["Nội dung", "Thông tin"],
    [
        ("Ngày báo cáo", "13/09/2026"),
        ("Phạm vi", "Admin Web, KPI tư vấn viên, API, PostgreSQL và đồng bộ Google Sheets"),
        ("Trạng thái", "Đã chạy xuyên suốt trên môi trường local và xác nhận dữ liệu trên Web"),
        ("Quyền truy cập Web", "Một tài khoản Admin duy nhất"),
        ("Dữ liệu trong báo cáo", "Không chứa tên thật, email thật hoặc thông tin định danh học sinh"),
    ],
    widths=[1.6, 5.2],
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(20)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Kết luận ngắn")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.left_indent = Inches(0.55)
p.paragraph_format.right_indent = Inches(0.55)
p.add_run(
    "Hệ thống đánh giá 4 KPI hiệu suất theo quy tắc đạt tối thiểu 3/4 và tách tải ca thành điều kiện an toàn bắt buộc; "
    "thiếu dữ liệu được tách riêng, chấm công HR không bị dùng để kết luận chuyên môn, Web chỉ cho phép đăng nhập bằng một tài khoản Admin cấu hình trên máy chủ, và dữ liệu học sinh từ hai tab THCS/THPT đã đồng bộ vào PostgreSQL local."
)

doc.add_page_break()

add_heading(doc, "1 Tóm tắt điều hành", 1)
doc.add_paragraph(
    "Mục tiêu của lần triển khai này là làm cho kết quả KPI dễ giải thích, có thể kiểm tra lại từ dữ liệu nguồn và phù hợp với phạm vi Admin Web của dự án. Phiên bản mới không coi mọi tiêu chí như nhau trong mọi tình huống và không tự động gắn nhãn “chưa đạt” khi dữ liệu chưa đủ."
)
add_bullet(doc, "Chuẩn hóa 4 KPI hiệu suất và 1 điều kiện an toàn tải ca với công thức, ngưỡng và bằng chứng đi kèm.")
add_bullet(doc, "Kết luận “Đạt” khi có ít nhất 3/4 KPI hiệu suất đạt và tải ca nằm trong giới hạn an toàn.")
add_bullet(doc, "Kết luận “Chưa đủ dữ liệu” khi tải ca bằng 0/chưa xác định hoặc có dưới 3 KPI hiệu suất đủ bằng chứng.")
add_bullet(doc, "Tách chấm công/giờ làm HR khỏi điểm chuyên môn; HR chỉ là tín hiệu cần rà soát vận hành.")
add_bullet(doc, "Dùng tối thiểu 5 phản hồi ẩn danh trước khi tính tiêu chí trải nghiệm học sinh.")
add_bullet(doc, "Khóa luồng đăng nhập Web ở một tài khoản Admin duy nhất; Counselor không thể đăng nhập Web.")
add_bullet(doc, "Đồng bộ hồ sơ học sinh từ hai tab Google Sheets THCS và THPT qua Apps Script; Web chỉ đọc và yêu cầu chỉnh sửa tại Sheet.")

add_heading(doc, "2 Phạm vi và căn cứ triển khai", 1)
doc.add_paragraph(
    "Phạm vi bám theo phân công WEB-01 đến WEB-06: đăng nhập quản trị, dashboard, danh sách/chi tiết tư vấn viên, biểu đồ và bộ lọc theo kỳ. SQL là nguồn dữ liệu chuẩn. Hồ sơ học sinh từ hai tab Google Sheets THCS và THPT đi qua Apps Script có khóa xác thực, sau đó được cập nhật vào PostgreSQL và hiển thị lại trên Web."
)
add_table(
    doc,
    ["Hạng mục", "Quyết định triển khai", "Ý nghĩa cho leader"],
    [
        ("Đăng nhập", "Một Admin cấu hình bằng biến môi trường", "Không phát sinh màn hình hoặc quyền đăng nhập Counselor trong MVP"),
        ("KPI", "4 KPI hiệu suất và 1 điều kiện an toàn, tính lại ở backend và frontend", "Không dùng ít ca để suy ra hiệu suất tốt"),
        ("Tổng hợp", "3/4 KPI hiệu suất + tải ca an toàn", "Cho phép một sai lệch nhỏ nhưng không bỏ qua rủi ro quá tải"),
        ("Thiếu dữ liệu", "Trạng thái riêng", "Không đánh đồng thiếu bằng chứng với năng lực kém"),
        ("HR", "Theo dõi riêng", "Dùng để điều phối nhân sự, không thay thế đánh giá chuyên môn"),
        ("Riêng tư", "Tổng hợp ẩn danh, ngưỡng mẫu", "Giảm nguy cơ suy ngược dữ liệu học sinh"),
        ("Quản lý Student", "Chỉnh sửa tại Google Sheets, Web chỉ đọc", "Phù hợp phạm vi MVP và chi phí hiện tại"),
    ],
    widths=[1.15, 2.65, 3.0],
)

doc.add_page_break()

add_heading(doc, "3 Mô hình KPI đã triển khai", 1)
doc.add_paragraph(
    "Mỗi tiêu chí có ngưỡng đạt rõ ràng. Tải ca là điều kiện an toàn, không được tính là thành tích và không đóng góp vào điểm hiệu suất. Điểm tổng hợp chỉ chuẩn hóa trên tổng trọng số 85% của bốn KPI hiệu suất. Giao diện ưu tiên nhãn dễ hiểu “Đạt / Cần cải thiện / Chưa đủ dữ liệu”."
)
add_table(
    doc,
    ["#", "Tiêu chí", "Công thức / nguồn", "Ngưỡng", "Trọng số"],
    [
        ("1", "Tải ca quy đổi theo FTE", "Tổng trọng số ca đang hoạt động ÷ FTE", "1–20: an toàn; 0: thiếu dữ liệu; >20: quá tải", "Không tính điểm"),
        ("2", "Thời gian dành cho học sinh", "Giờ phiên tư vấn hoàn thành ÷ giờ hỗ trợ đã đăng ký", "≥ 80%", "20%"),
        ("3", "Hoàn thành lịch tư vấn đủ điều kiện", "Phiên hoàn thành ÷ phiên đến hạn đủ điều kiện", "≥ 80%", "25%"),
        ("4", "Theo dõi bài đánh giá", "Bài hoàn thành ÷ bài được giao đủ điều kiện", "≥ 80%", "20%"),
        ("5", "Trải nghiệm/kết quả cảm nhận", "Điểm phản hồi trung bình × 20", "≥ 80% (tương đương 4/5) và ≥ 5 phản hồi", "20%"),
    ],
    widths=[0.35, 1.75, 2.65, 1.55, 0.6],
)

add_heading(doc, "4 Quy tắc kết luận tổng thể", 1)
add_table(
    doc,
    ["Trạng thái", "Điều kiện", "Cách hiểu"],
    [
        ("ĐẠT", "Ít nhất 3/4 KPI hiệu suất đủ dữ liệu và đạt; tải ca từ 1 đến 20 ca quy đổi/FTE", "Đáp ứng chuẩn trong kỳ"),
        ("CHƯA ĐẠT", "Có ít nhất 3 KPI hiệu suất đủ dữ liệu nhưng dưới 3 KPI đạt, hoặc tải ca vượt 20", "Cần rà soát nguyên nhân và kế hoạch cải thiện"),
        ("CHƯA ĐỦ DỮ LIỆU", "Tải ca bằng 0/chưa xác định, dưới 3 KPI hiệu suất đủ bằng chứng, hoặc bộ KPI sai cấu trúc", "Chưa được phép kết luận hiệu suất"),
    ],
    widths=[1.35, 3.25, 2.3],
)
doc.add_paragraph(
    "Ví dụ: tư vấn viên có tải ca an toàn và 3 KPI hiệu suất đạt vẫn có thể được kết luận “Đạt”, kể cả KPI còn lại chưa đạt hoặc chưa đủ bằng chứng. Nếu tải ca bằng 0, hoặc chỉ có 2 KPI hiệu suất đủ dữ liệu, kết quả bắt buộc là “Chưa đủ dữ liệu”."
)

doc.add_page_break()

add_heading(doc, "5 Chi tiết công thức và kiểm soát sai lệch", 1)
add_heading(doc, "5 1 Tải ca quy đổi theo FTE", 2)
doc.add_paragraph(
    "Tải ca không còn chỉ đếm số học sinh. Mỗi ca có hệ số độ phức tạp (mặc định 1,0; đề xuất 1,5 cho ca cần theo dõi nhiều; 2–3 cho ca nguy cơ cao/khủng hoảng) và được chia theo tỷ lệ FTE của tư vấn viên. Từ 1 đến 20 là trong giới hạn an toàn; trên 20 là quá tải. Giá trị 0 được ghi là “Chưa đủ dữ liệu”, không phải “Đạt”."
)
add_heading(doc, "5 2 Thời gian dịch vụ cho học sinh", 2)
doc.add_paragraph(
    "Tỷ lệ được tính từ thời lượng các phiên đã hoàn thành trên tổng giờ hỗ trợ đã đăng ký. Khi mẫu số bằng 0, hệ thống trả về “chưa đủ dữ liệu” thay vì chia cho 0 hoặc gán 0%. Hiện chưa tính công việc hỗ trợ gián tiếp vì cơ sở dữ liệu chưa có nhật ký thời gian có cấu trúc."
)
add_heading(doc, "5 3 Hoàn thành lịch tư vấn đủ điều kiện", 2)
doc.add_paragraph(
    "Chỉ tính lịch đã đến hạn trong kỳ. Loại trừ lịch tương lai, học sinh hủy, học sinh vắng và lịch được đổi. Chỉ đánh giá khi có ít nhất 5 phiên đủ điều kiện; dưới ngưỡng này hiển thị “Chưa đủ lịch”."
)
add_heading(doc, "5 4 Theo dõi bài đánh giá", 2)
doc.add_paragraph(
    "Chỉ tính bài được giao còn hiệu lực. Các trạng thái từ chối, rút lui, không còn cần hoặc đã hủy được loại khỏi mẫu số. Bài có kết quả hoặc lượt làm hoàn thành được ghi nhận là hoàn thành. Chỉ đánh giá khi có ít nhất 3 bài được giao đủ điều kiện."
)
add_heading(doc, "5 5 Trải nghiệm và kết quả cảm nhận", 2)
doc.add_paragraph(
    "Điểm trung bình phản hồi 1–5 được quy đổi sang thang phần trăm. Chỉ đánh giá khi có ít nhất 5 phản hồi ẩn danh trong kỳ. Nếu ít hơn, giao diện hiển thị “Chưa đủ phản hồi” và không coi đây là một tiêu chí thất bại."
)

add_heading(doc, "6 Chấm công HR và quyền riêng tư", 1)
doc.add_paragraph(
    "Dữ liệu ngày làm, tổng giờ đăng ký, số ngày vượt 8 giờ và số tuần làm hơn 6 ngày được hiển thị trong một khu vực riêng. Các chỉ số này hỗ trợ điều phối và an toàn lao động, nhưng không cộng hoặc trừ vào kết luận chuyên môn."
)
add_bullet(doc, "Không đưa nội dung phản hồi thô hoặc thông tin nhận diện học sinh vào KPI.")
add_bullet(doc, "Dữ liệu phân tích phản hồi có ngưỡng mẫu tối thiểu trước khi hiển thị chi tiết.")
add_bullet(doc, "Báo cáo và giao diện dùng số tổng hợp; dữ liệu cá nhân chỉ nằm trong hệ thống nguồn có kiểm soát.")

add_heading(doc, "7 Mô hình một tài khoản Admin", 1)
doc.add_paragraph(
    "Web hiện chỉ chấp nhận đúng tài khoản Admin được khai báo bằng ADMIN_EMAIL và ADMIN_PASSWORD_HASH ở backend. Backend không còn tra cứu tài khoản Counselor khi đăng nhập; frontend xóa phiên cũ không phải Admin và từ chối phản hồi đăng nhập có vai trò khác Admin."
)
add_table(
    doc,
    ["Lớp kiểm soát", "Đã triển khai"],
    [
        ("Backend", "Chỉ so khớp tài khoản Admin cấu hình; JWT phát hành với vai trò admin"),
        ("Frontend", "Danh sách demo chỉ còn một Admin; phiên lưu cũ có role Counselor bị loại bỏ"),
        ("Điều hướng", "Đăng nhập thành công luôn vào dashboard Admin"),
        ("Thông báo", "Nếu API trả về vai trò không phải Admin, giao diện báo không có quyền truy cập"),
    ],
    widths=[1.6, 5.2],
)

p = doc.add_paragraph()
r = p.add_run("Lưu ý quản trị quan trọng: ")
r.bold = True
p.add_run(
    "một tài khoản dùng chung đáp ứng yêu cầu MVP nhưng làm giảm khả năng truy vết cá nhân trong audit log. Nên lưu mật khẩu trong password manager của đơn vị, giới hạn người biết mật khẩu, thay đổi định kỳ và chuyển sang tài khoản định danh/MFA khi dự án qua giai đoạn thử nghiệm."
)

add_heading(doc, "8 Thay đổi kỹ thuật chính", 1)
add_table(
    doc,
    ["Khu vực", "Nội dung thay đổi", "Kết quả"],
    [
        ("Frontend – chính sách KPI", "4 KPI hiệu suất, quy tắc 3/4, điều kiện an toàn tải ca và trạng thái thiếu dữ liệu", "Giao diện và dữ liệu demo dùng cùng một quy tắc"),
        ("Frontend – màn hình", "Dashboard, danh sách, chi tiết, biểu đồ và đăng nhập", "Nhãn tiếng Việt rõ ràng, phân biệt đạt/chưa đạt/thiếu dữ liệu"),
        ("Backend – tính toán", "Công thức KPI, bằng chứng, trọng số, HR riêng", "API trả dữ liệu có thể kiểm tra lại"),
        ("Backend – truy vấn", "FTE, trọng số ca, giờ đăng ký, phiên đủ điều kiện, bài đánh giá và phản hồi", "Giảm sai lệch do mẫu số không phù hợp"),
        ("Backend – xác thực", "Loại bỏ nhánh đăng nhập Counselor", "Web chỉ còn một principal Admin"),
        ("Google Sheets", "Apps Script đồng bộ hai tab THCS/THPT; dropdown trạng thái; bỏ qua dòng trống; xác thực Bearer", "Upsert theo số điện thoại/email; Ngừng theo dõi chuyển hồ sơ thành INACTIVE"),
        ("Database", "Migration bổ sung fte_ratio, case_weight và school_level", "Giữ tương thích dữ liệu cũ và tách danh sách học sinh theo cấp"),
    ],
    widths=[1.45, 3.0, 2.35],
)

add_heading(doc, "9 Kết quả kiểm thử", 1)
add_table(
    doc,
    ["Hạng mục", "Kết quả", "Ghi chú"],
    [
        ("Frontend type-check", "ĐẠT", "TypeScript không có lỗi"),
        ("Frontend KPI policy", "ĐẠT", "Kiểm tra 3/4, tải ca an toàn, mức bằng chứng tối thiểu và dữ liệu demo"),
        ("Frontend production build", "ĐẠT", "Build thành công; có cảnh báo bundle >500 kB, không chặn chạy"),
        ("Frontend smoke test", "ĐẠT", "Máy chủ cục bộ trả HTTP 200"),
        ("Backend type-check", "ĐẠT", "TypeScript không có lỗi"),
        ("Backend automated tests", "34/34 ĐẠT", "Bao gồm KPI, phân quyền, bảo mật mẫu nhỏ, CRUD, soft delete từ Sheet và SQL"),
        ("Backend production build", "ĐẠT", "Biên dịch thư mục dist thành công"),
        ("Đồng bộ Google Sheets", "ĐẠT", "Apps Script đồng bộ 2 dòng hợp lệ, bỏ qua 5 dòng trống; trigger đổi trạng thái chạy thành công"),
        ("Đối chiếu Web và DB", "ĐẠT", "15 học sinh: 8 THCS và 7 THPT ở cả PostgreSQL và Web"),
    ],
    widths=[2.05, 1.0, 3.75],
)

add_heading(doc, "10 Hướng dẫn đưa lên môi trường tích hợp", 1)
add_number(doc, "Sao lưu cơ sở dữ liệu và xác nhận phiên bản schema hiện tại.")
add_number(doc, "Chạy migration database/migrations/001_professional_kpi_inputs.sql và 002_student_school_level.sql đúng một lần.")
add_number(doc, "Cấu hình DATABASE_URL, JWT_SECRET, ADMIN_EMAIL, ADMIN_PASSWORD_HASH, GOOGLE_SHEETS_SYNC_SECRET và CORS_ORIGINS; không đưa mật khẩu thô vào mã nguồn.")
add_number(doc, "Build và khởi động backend; kiểm tra health endpoint và đăng nhập bằng tài khoản Admin.")
add_number(doc, "Cấu hình VITE_API_BASE_URL ở frontend, build bản production và kiểm tra dashboard theo ba kỳ báo cáo.")
add_number(doc, "Đối chiếu ngẫu nhiên ít nhất ba tư vấn viên với SQL nguồn trước khi nghiệm thu.")
add_number(doc, "Cấu hình BACKEND_BASE_URL và GOOGLE_SHEETS_SYNC_SECRET trong Apps Script, cài trigger on form submit/on edit và chạy syncAllStudents một lần.")
add_number(doc, "Dùng URL HTTPS ổn định ở môi trường tích hợp; tunnel localhost.run chỉ phục vụ chạy thử local.")

add_heading(doc, "11 Điều kiện nghiệm thu đề xuất", 1)
add_bullet(doc, "Leader xác nhận quy tắc 3/4 KPI hiệu suất và tải ca an toàn từ 1 đến 20 ca quy đổi/FTE.")
add_bullet(doc, "Nghiệp vụ xác nhận bảng hệ số case_weight và quy trình cập nhật FTE.")
add_bullet(doc, "Tài khoản Admin thật được tạo bằng bcrypt hash, không dùng mật khẩu demo.")
add_bullet(doc, "Kết quả “Chưa đủ dữ liệu” xuất hiện đúng khi không đủ mẫu hoặc mẫu số bằng 0.")
add_bullet(doc, "Chấm công HR không làm thay đổi kết luận KPI chuyên môn.")
add_bullet(doc, "Counselor không thể đăng nhập Web dù có hồ sơ hoặc bản ghi account trong database.")
add_bullet(doc, "Thêm hoặc sửa hồ sơ hợp lệ trên tab THCS/THPT cập nhật database; Web hiển thị lại sau khi bấm Làm mới hoặc tải lại trang.")
add_bullet(doc, "Chọn Ngừng theo dõi trên Sheet chuyển hồ sơ thành INACTIVE và ẩn khỏi Web; chọn lại Đang hoạt động để khôi phục.")

doc.add_page_break()

add_heading(doc, "12 Giới hạn hiện tại và đề xuất giai đoạn tiếp theo", 1)
add_table(
    doc,
    ["Giới hạn", "Ảnh hưởng", "Đề xuất"],
    [
        ("Chưa có metadata chiều điểm của từng thang đo", "Không thể kết luận tiến bộ lâm sàng từ thay đổi score một cách an toàn", "Bổ sung direction, baseline window và clinically_significant_change theo từng test"),
        ("Chưa có nhật ký thời gian hỗ trợ gián tiếp", "KPI thời gian chỉ tính giờ phiên trực tiếp", "Thêm activity log có loại hoạt động và thời lượng"),
        ("Một tài khoản Admin dùng chung", "Audit không phân biệt người thao tác", "Giữ cho MVP; sau đó chuyển sang tài khoản cá nhân + MFA"),
        ("Bundle frontend còn lớn", "Tải trang có thể chậm trên thiết bị yếu", "Tách trang analytics bằng lazy loading"),
        ("Tunnel chạy thử là URL tạm thời", "URL có thể đổi khi kết nối lại và làm trigger ngừng đồng bộ", "Triển khai backend trên HTTPS ổn định trước nghiệm thu"),
        ("Không hỗ trợ xóa vật lý từ Sheet", "Dữ liệu lịch sử tiếp tục được giữ trong database", "Dùng dropdown Ngừng theo dõi để soft delete; xóa vật lý là quy trình quản trị riêng có sao lưu"),
        ("Web không tự đẩy cập nhật tức thời", "Admin cần bấm Làm mới hoặc tải lại trang", "Giữ cho MVP; cân nhắc polling hoặc realtime ở giai đoạn sau"),
        ("Ngưỡng là cấu hình nghiệp vụ ban đầu", "Có thể chưa phản ánh đủ điều kiện địa phương", "Chạy pilot 1–2 kỳ và hiệu chỉnh bằng hội đồng chuyên môn"),
    ],
    widths=[1.75, 2.35, 2.7],
)

add_heading(doc, "13 Truy vết yêu cầu", 1)
add_table(
    doc,
    ["Mã", "Nội dung liên quan", "Trạng thái sau triển khai"],
    [
        ("WEB-01", "Đăng nhập quản trị", "Một tài khoản Admin; Counselor bị chặn"),
        ("WEB-02", "Dashboard tổng quan", "Hiển thị kết quả theo quy tắc mới"),
        ("WEB-03", "Danh sách hiệu suất tư vấn viên", "Bộ lọc và ba trạng thái kết luận"),
        ("WEB-04", "Biểu đồ KPI", "Nhãn 3/4 KPI hiệu suất + điều kiện an toàn, nhóm cần rà soát"),
        ("WEB-05", "Chi tiết tư vấn viên", "5 thẻ KPI, công thức, bằng chứng và HR riêng"),
        ("WEB-06", "Bộ lọc kỳ báo cáo", "Giữ ba kỳ: tháng này, tháng trước, toàn thời gian"),
        ("Student", "Danh sách THCS/THPT", "Web chỉ đọc; chỉnh sửa và soft delete bằng trạng thái tại Sheet; đồng bộ qua backend"),
    ],
    widths=[0.75, 2.75, 3.3],
)

add_heading(doc, "14 Tài liệu tham chiếu", 1)
refs = [
    "Danh_ARLFreshers_Project2_Summary/Task_Breakdown_Trung_Tam_Tu_Van_Tam_Ly.docx",
    "Danh_ARLFreshers_Project2_Summary/Requirements/Requirements_Specification_Trung_Tam_Tu_Van_Tam_Ly.docx",
    "ASCA – The School Counselor and Annual Performance Appraisal: https://www.schoolcounselor.org/Standards-Positions/Position-Statements/ASCA-Position-Statements/The-School-Counselor-and-Annual-Performance-Apprai",
    "ASCA – The Professional Counselor and Use of Support Staff: https://schoolcounselor.org/Standards-Positions/Position-Statements/ASCA-Position-Statements/The-Professional-Counselor-and-Use-of-Support-Staf",
    "ASCA Ethical Standards for School Counselors: https://www.schoolcounselor.org/About-School-Counseling/Ethical-Responsibilities/ASCA-Ethical-Standards-for-School-Counselors-%281%29",
]
for ref in refs:
    add_bullet(doc, ref)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— HẾT BÁO CÁO —")
r.bold = True
r.font.size = Pt(10)
r.font.color.rgb = RGBColor.from_string(BLUE)

doc.core_properties.title = "Báo cáo triển khai KPI và cổng Web một tài khoản Admin"
doc.core_properties.subject = "Bàn giao triển khai Admin Web – Digital Twin"
doc.core_properties.author = "Nhóm dự án Digital Twin"
doc.core_properties.keywords = "KPI, Admin Web, Digital Twin, tư vấn tâm lý học đường"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
